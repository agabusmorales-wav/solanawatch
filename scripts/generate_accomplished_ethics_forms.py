import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PROJECT_DIR = r"C:\Users\agabu\Downloads\CODE PROJECTS\SOLANAWATCH"
DOWNLOADS_DIR = r"C:\Users\agabu\Downloads"
DEST_DIRS = [PROJECT_DIR, DOWNLOADS_DIR]

LOGO_DIR = r"C:\Users\agabu\.gemini\antigravity-cli\brain\e1ec8bb1-0a64-46ec-80e7-07d55ccadc95\scratch\logos"
UST_LOGO = os.path.join(LOGO_DIR, "logo_1.png")
CICS_LOGO = os.path.join(LOGO_DIR, "logo_2.png")
REC_LOGO = os.path.join(LOGO_DIR, "logo_3.png")

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="000000", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement('w:cantSplit'))

def set_header(doc, form_id, form_name, date_str="6 June 2024"):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(13.0)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    footer = section.footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f_p.paragraph_format.space_before = Pt(0)
    f_p.paragraph_format.space_after = Pt(0)
    run = f_p.add_run(f"{form_id}: {form_name}\n{date_str}   |   Page ")
    run.font.name = "Arial"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    footer.paragraphs[0]._p.append(fldSimple)

def add_header_table(doc, form_title_lines):
    tbl = doc.add_table(rows=1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    tbl.columns[0].width = Inches(1.1)
    tbl.columns[1].width = Inches(4.9)
    tbl.columns[2].width = Inches(1.1)

    c0 = tbl.cell(0, 0)
    c1 = tbl.cell(0, 1)
    c2 = tbl.cell(0, 2)

    c0.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    c1.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    c2.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if os.path.exists(UST_LOGO):
        r0 = p0.add_run()
        r0.add_picture(UST_LOGO, width=Inches(0.95))

    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(2)
    p1.paragraph_format.line_spacing = 1.15

    r_inst = p1.add_run("UNIVERSITY OF SANTO TOMAS\n")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(11)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(0x11, 0x11, 0x11)

    r_coll = p1.add_run("COLLEGE OF INFORMATION AND COMPUTING SCIENCES\n")
    r_coll.font.name = "Arial"
    r_coll.font.size = Pt(10)
    r_coll.font.bold = True
    r_coll.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    for line, is_bold, sz, col in form_title_lines:
        r_line = p1.add_run(line + "\n")
        r_line.font.name = "Arial"
        r_line.font.size = Pt(sz)
        r_line.font.bold = is_bold
        r_line.font.color.rgb = col

    p2 = c2.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if os.path.exists(CICS_LOGO):
        r2 = p2.add_run()
        r2.add_picture(CICS_LOGO, width=Inches(0.95))

    tblPr = tbl._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="none"/>\n'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>\n'
        f'  <w:left w:val="none"/>\n'
        f'  <w:right w:val="none"/>\n'
        f'  <w:insideH w:val="none"/>\n'
        f'  <w:insideV w:val="none"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

print("Unified ethics forms generator ready.")
